# app.py
# Sondage pour récolter des questions sur le Togo.
# Lancement : python app.py  puis ouvrir http://localhost:8000

import csv
import io
import os
import threading
import unicodedata
import uuid
from datetime import date, datetime

from flask import Flask, jsonify, request, send_file, send_from_directory
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from werkzeug.utils import secure_filename

DOSSIER = os.path.dirname(os.path.abspath(__file__))
FICHIER_REPONSES = os.path.join(DOSSIER, "reponses_sondage.csv")
# Fichiers joints : jamais servis par le site, à télécharger depuis PythonAnywhere (onglet Files)
DOSSIER_FICHIERS = os.path.join(DOSSIER, "fichiers_recus")
INDEX_FICHIERS = os.path.join(DOSSIER_FICHIERS, "liste_fichiers.csv")

COLONNES = ["Catégorie", "Question", "Option 1", "Option 2", "Option 3", "Réponse Correcte", "Proposé par", "Date"]
COLONNES_FICHIERS = ["Date", "Proposé par", "Catégorie", "Fichier enregistré", "Nom d'origine"]
CATEGORIES = {"Histoire", "Géographie", "Culture", "Sport", "Institutions"}

LONGUEUR_MAX_QUESTION = 200
LONGUEUR_MAX_REPONSE = 80
LONGUEUR_MAX_PSEUDO = 40

# Fichiers acceptés : documents, images, audio et vidéo (pas de programmes exécutables)
EXTENSIONS_AUTORISEES = {
    "pdf", "doc", "docx", "odt", "rtf", "txt", "csv", "xls", "xlsx", "ods", "ppt", "pptx", "odp",
    "jpg", "jpeg", "png", "gif", "webp", "heic",
    "mp3", "m4a", "ogg", "opus", "wav", "aac",
    "mp4", "mov", "3gp", "webm",
}
TAILLE_MAX_FICHIER = 10 * 1024 * 1024

# Modèle à remplir pour envoyer beaucoup de questions d'un coup
COLONNES_MODELE = ["Catégorie", "Question", "Réponse 1", "Réponse 2", "Réponse 3", "Bonne réponse"]
EXEMPLE = ["Histoire", "En quelle année le Togo est-il devenu indépendant ?", "1958", "1960", "1963", "1960"]
LIGNES_MAX_IMPORT = 500
# Le compte gratuit PythonAnywhere a 512 Mo : on garde de la marge
TAILLE_MAX_DOSSIER = 300 * 1024 * 1024

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = TAILLE_MAX_FICHIER + 100 * 1024
verrou = threading.Lock()


def lire_reponses():
    if not os.path.exists(FICHIER_REPONSES):
        return []
    with open(FICHIER_REPONSES, encoding="utf-8", newline="") as f:
        return [ligne for ligne in csv.DictReader(f) if ligne.get("Question")]


def ajouter_ligne(chemin, colonnes, ligne):
    nouveau_fichier = not os.path.exists(chemin)
    with open(chemin, "a", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        if nouveau_fichier:
            writer.writerow(colonnes)
        writer.writerow(ligne)


def taille_dossier_fichiers():
    if not os.path.isdir(DOSSIER_FICHIERS):
        return 0
    return sum(e.stat().st_size for e in os.scandir(DOSSIER_FICHIERS) if e.is_file())


def nettoyer(texte, longueur_max):
    texte = " ".join(str(texte or "").split())
    # Empêche Excel d'interpréter le texte comme une formule
    texte = texte.lstrip("=+-@ ")
    return texte[:longueur_max]


def erreur(texte, code=400):
    return jsonify(ok=False, erreur=texte), code


def sans_accents(texte):
    texte = unicodedata.normalize("NFD", str(texte or "").strip().lower())
    return "".join(c for c in texte if unicodedata.category(c) != "Mn")


CATEGORIES_SANS_ACCENTS = {sans_accents(c): c for c in CATEGORIES}


def verifier_question(categorie, question, options, index_bonne):
    """Renvoie le message d'erreur, ou None si la question est valable."""
    if categorie not in CATEGORIES:
        return "Choisissez une catégorie."
    if len(question) < 10:
        return "La question est trop courte."
    if len(options) != 3 or not all(options):
        return "Remplissez les trois réponses."
    if len({o.lower() for o in options}) != 3:
        return "Les trois réponses doivent être différentes."
    if index_bonne not in (0, 1, 2):
        return "Cochez la bonne réponse."
    return None


def texte_cellule(valeur):
    # Excel transforme « 1960 » en 1960.0 et certaines réponses en dates
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    elif isinstance(valeur, datetime):
        valeur = valeur.date().isoformat()
    return "" if valeur is None else str(valeur)


def lire_tableau(fichier, extension):
    """Lit les lignes d'un fichier .xlsx ou .csv envoyé avec le modèle."""
    if extension == "xlsx":
        classeur = load_workbook(fichier, read_only=True, data_only=True)
        feuille = classeur["Questions"] if "Questions" in classeur.sheetnames else classeur.worksheets[0]
        lignes = [[texte_cellule(v) for v in ligne]
                  for ligne in feuille.iter_rows(max_row=LIGNES_MAX_IMPORT + 1, max_col=6, values_only=True)]
        classeur.close()
        return lignes

    brut = fichier.read()
    try:
        texte = brut.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel en français enregistre souvent les CSV en Windows-1252
        texte = brut.decode("cp1252", errors="replace")
    try:
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=",;\t")
    except csv.Error:
        dialecte = csv.excel
    return list(csv.reader(io.StringIO(texte), dialecte))[:LIGNES_MAX_IMPORT + 1]


def ligne_vers_question(ligne):
    """Transforme une ligne du modèle en (catégorie, question, options, index_bonne)."""
    ligne = (list(ligne) + [""] * 6)[:6]
    categorie = CATEGORIES_SANS_ACCENTS.get(sans_accents(ligne[0]), nettoyer(ligne[0], 40))
    question = nettoyer(ligne[1], LONGUEUR_MAX_QUESTION)
    options = [nettoyer(o, LONGUEUR_MAX_REPONSE) for o in ligne[2:5]]
    bonne = nettoyer(ligne[5], LONGUEUR_MAX_REPONSE)

    # La bonne réponse peut être écrite en entier ou donnée par son numéro (1, 2 ou 3)
    if bonne in ("1", "2", "3") and bonne not in options:
        index_bonne = int(bonne) - 1
    else:
        index_bonne = next((i for i, o in enumerate(options) if o and o.lower() == bonne.lower()), -1)
    return categorie, question, options, index_bonne


@app.errorhandler(413)
def fichier_trop_gros(_):
    return erreur("Le fichier est trop lourd (10 Mo maximum).", 413)


@app.route("/")
def accueil():
    return send_from_directory(DOSSIER, "index.html")


@app.route("/envoyer", methods=["POST"])
def envoyer():
    donnees = request.form

    # Champ piège invisible : seuls les robots le remplissent
    if donnees.get("site_web"):
        return jsonify(ok=True)

    categorie = nettoyer(donnees.get("categorie"), 40)
    question = nettoyer(donnees.get("question"), LONGUEUR_MAX_QUESTION)
    options = [nettoyer(o, LONGUEUR_MAX_REPONSE) for o in donnees.getlist("options")][:3]
    pseudo = nettoyer(donnees.get("pseudo"), LONGUEUR_MAX_PSEUDO) or "Anonyme"
    fichier = request.files.get("fichier")
    if fichier and not fichier.filename:
        fichier = None

    # Une question est en cours de saisie dès qu'un de ses champs est rempli
    avec_question = bool(question or any(options))

    if not avec_question and not fichier:
        return erreur("Écrivez une question ou joignez un fichier.")
    if categorie not in CATEGORIES:
        return erreur("Choisissez une catégorie.")

    if avec_question:
        try:
            index_bonne = int(donnees.get("bonne_reponse"))
        except (TypeError, ValueError):
            index_bonne = -1

        probleme = verifier_question(categorie, question, options, index_bonne)
        if probleme:
            return erreur(probleme)

    if fichier:
        extension = fichier.filename.rsplit(".", 1)[-1].lower() if "." in fichier.filename else ""
        if extension not in EXTENSIONS_AUTORISEES:
            return erreur("Ce type de fichier n'est pas accepté (photos, documents, audio et vidéo uniquement).")

    with verrou:
        if avec_question and any(q["Question"].strip().lower() == question.lower() for q in lire_reponses()):
            return erreur("Cette question a déjà été proposée.")

        if fichier:
            if taille_dossier_fichiers() > TAILLE_MAX_DOSSIER:
                return erreur("Nous ne pouvons plus recevoir de fichiers pour le moment. Écrivez votre question dans le formulaire.", 507)

            os.makedirs(DOSSIER_FICHIERS, exist_ok=True)
            nom_propre = secure_filename(fichier.filename) or f"fichier.{extension}"
            nom_enregistre = f"{date.today().isoformat()}_{uuid.uuid4().hex[:8]}_{nom_propre}"
            fichier.save(os.path.join(DOSSIER_FICHIERS, nom_enregistre))
            ajouter_ligne(INDEX_FICHIERS, COLONNES_FICHIERS,
                          [date.today().isoformat(), pseudo, categorie, nom_enregistre, nettoyer(fichier.filename, 150)])

        if avec_question:
            ajouter_ligne(FICHIER_REPONSES, COLONNES,
                          [categorie, question, *options, options[index_bonne], pseudo, date.today().isoformat()])

    return jsonify(ok=True)


@app.route("/modele.xlsx")
def modele_excel():
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Questions"
    feuille.append(COLONNES_MODELE)
    feuille.append(EXEMPLE)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for cellule in feuille[2]:
        cellule.font = Font(italic=True, color="808080")
    for colonne, largeur in zip("ABCDEF", (16, 60, 25, 25, 25, 25)):
        feuille.column_dimensions[colonne].width = largeur
    feuille.freeze_panes = "A2"

    # Liste déroulante pour choisir la catégorie
    liste = DataValidation(type="list", formula1='"' + ",".join(sorted(CATEGORIES)) + '"', allow_blank=True)
    liste.error = "Choisissez une catégorie dans la liste."
    feuille.add_data_validation(liste)
    liste.add(f"A2:A{LIGNES_MAX_IMPORT + 1}")

    aide = classeur.create_sheet("Mode d'emploi")
    for texte in [
        "Une question par ligne, dans la feuille « Questions ».",
        "Catégorie : Histoire, Géographie, Culture, Sport ou Institutions.",
        "Trois réponses différentes, dont une seule est juste.",
        "Bonne réponse : recopier la bonne réponse, ou écrire son numéro (1, 2 ou 3).",
        "La ligne grise d'exemple est ignorée : vous pouvez la laisser ou l'effacer.",
        f"{LIGNES_MAX_IMPORT} questions maximum par fichier.",
        "Une fois rempli, envoyez le fichier avec le bouton « Importer » du formulaire.",
    ]:
        aide.append([texte])
    aide.column_dimensions["A"].width = 90

    contenu = io.BytesIO()
    classeur.save(contenu)
    contenu.seek(0)
    return send_file(contenu, as_attachment=True, download_name="modele_questions_togo.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/modele.csv")
def modele_csv():
    contenu = io.StringIO()
    writer = csv.writer(contenu)
    writer.writerow(COLONNES_MODELE)
    writer.writerow(EXEMPLE)
    # BOM pour qu'Excel affiche correctement les accents
    donnees = io.BytesIO(("﻿" + contenu.getvalue()).encode("utf-8"))
    return send_file(donnees, as_attachment=True, download_name="modele_questions_togo.csv", mimetype="text/csv")


@app.route("/importer", methods=["POST"])
def importer():
    if request.form.get("site_web"):
        return jsonify(ok=True, ajoutees=0, erreurs=[])

    fichier = request.files.get("tableau")
    if not fichier or not fichier.filename:
        return erreur("Choisissez le fichier rempli (.xlsx ou .csv).")
    extension = fichier.filename.rsplit(".", 1)[-1].lower() if "." in fichier.filename else ""
    if extension not in ("xlsx", "csv"):
        return erreur("Utilisez le modèle au format Excel (.xlsx) ou CSV.")
    pseudo = nettoyer(request.form.get("pseudo"), LONGUEUR_MAX_PSEUDO) or "Anonyme"

    try:
        lignes = lire_tableau(fichier, extension)
    except Exception:
        return erreur("Impossible de lire ce fichier. Partez du modèle à télécharger.")

    # La première ligne est l'en-tête si elle commence par « Catégorie »
    debut = 1
    if lignes and sans_accents(lignes[0][0] if lignes[0] else "") == "categorie":
        lignes, debut = lignes[1:], 2
    if len(lignes) > LIGNES_MAX_IMPORT:
        return erreur(f"Le fichier contient trop de lignes ({LIGNES_MAX_IMPORT} questions maximum).")

    erreurs, nouvelles = [], []
    with verrou:
        deja_vues = {q["Question"].strip().lower() for q in lire_reponses()}
        for numero, ligne in enumerate(lignes, start=debut):
            if not any(str(c).strip() for c in ligne):
                continue
            categorie, question, options, index_bonne = ligne_vers_question(ligne)
            if question == EXEMPLE[1]:
                continue
            probleme = verifier_question(categorie, question, options, index_bonne)
            if not probleme and question.lower() in deja_vues:
                probleme = "Cette question a déjà été proposée."
            if probleme:
                erreurs.append(f"Ligne {numero} : {probleme}")
                continue
            deja_vues.add(question.lower())
            nouvelles.append([categorie, question, *options, options[index_bonne], pseudo, date.today().isoformat()])

        for nouvelle in nouvelles:
            ajouter_ligne(FICHIER_REPONSES, COLONNES, nouvelle)

    if not nouvelles and not erreurs:
        return erreur("Le fichier ne contient aucune question.")
    return jsonify(ok=True, ajoutees=len(nouvelles), erreurs=erreurs)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 8000)))
